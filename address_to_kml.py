'''
Reads in an Excel .xlsx spreadsheet, gets the addresses, looks them up
using googlemaps python API, and makes a KML file with the addresses
and lat/lon for viewing in Google Earth.
'''

from time import sleep
import argparse
from lxml import etree
from openpyxl import load_workbook
import googlemaps
from pykml.factory import KML_ElementMaker as KML

isNone = lambda x: x == None

class Contact_t:
    '''
    Object to represent the contact information for a customer.
    INPUT
    - firstName = customer's first name
    - lastName = customer's last name
    - emailAddress = customer's email address
    - phoneNumber = customer's phone number (as a string)
    - address = customer's street address (no city, state, or zip code)
    - city = name of customer's city
    - state = 2-letter abbreviation of customer's state
    - zipCode = customer's zip code (as a string)
    - lat = latitude (degrees N) of customer's address (float)
    - lon = longitude (degrees E) of customer's address (float)
    '''
    def __init__(self, firstName='MISSING', lastName='MISSING',
            emailAddress='MISSING', phoneNumber='(###) ###-####',
            address='MISSING', city='MISSING', state='MISSING',
            zipCode='#####', lat=-999.0, lon=-999.0):
        self.firstName = firstName
        self.lastName = lastName
        self.emailAddress = emailAddress
        self.phoneNumber = phoneNumber
        self.address = address
        self.city = city
        self.state = state
        self.zipCode = zipCode
        self.lat = lat
        self.lon = lon
    
    def printAddressOneLine(self):
        '''
        Returns the mailing address on one line in the format:
        address, city, state zipCode
        '''
        return f'{self.address}, {self.city}, {self.state} {self.zipCode}'
        
    def printAddressLabel(self):
        '''
        Returns the mailing address as it would appear on an address label.
        '''
        return f'{self.firstName} {self.lastName}\n' \
            + f'{self.address}\n{self.city}, {self.state} {self.zipCode}'


class MulchContact_t(Contact_t):
    '''
    Object to contain contact and order information for a mulch customer.
    INPUT
    - firstName = customer's first name
    - lastName = customer's last name
    - emailAddress = customer's email address
    - phoneNumber = customer's phone number (as a string)
    - address = customer's street address (no city, state, or zip code)
    - city = name of customer's city
    - state = 2-letter abbreviation of customer's state
    - zipCode = customer's zip code (as a string)
    - nBags = number of bags of mulch ordered (int)
    - notes = delivery notes (string)
    - phoneNumber2 = customer's secondary phone number (as a string)
    '''
    def __init__(self, firstName='MISSING', lastName='MISSING',
            emailAddress='MISSING', phoneNumber='(###) ###-####',
            address='MISSING', city='MISSING', state='MISSING',
            zipCode='#####', nBags=0, notes='NONE',
            phoneNumber2='(###) ###-####'):
        super().__init__(firstName,lastName,emailAddress,phoneNumber,
            address,city,state,zipCode)
        self.nBags = nBags
        self.notes = notes
        self.phoneNumber2 = phoneNumber2

    @staticmethod
    def read_xlsx(xlsxFileName, sheetName, nRowsMax):
        '''
        Reads in a mulch xlsx file and return a list of MulchContact_t
        objects. 
        INPUT
        - xlsxFileName = path to Microsoft Excel (.xlsx) spreadsheet
            file containing mulch order information. The columns of the
            spreadsheet are as follows:
            The spreadsheet's first row is column headings.
            column 1 = first name
            column 2 = last name
            column 3 = address
            column 4 = city
            column 5 = state
            column 6 = zip code (numeric)
            column 7 = phone number
            column 8 = phone number (alternate)
            column 9 = e-mail address
            column 10 = number of bags
            column 11 = delivery (yes/no)
            column 12 = amount paid
            column 13 = notes
        - sheetName = name of workbook (spreadsheet) sheet / tab to use
        - nRowsMax = number of rows of spreadsheet to process (including
            the first row)
        OUTPUT
        Returns a list of MulchContact_t objects.
        '''
        outList = []
    
        # open xlsx file
        wb = load_workbook(filename=xlsxFileName)
        ws = wb[sheetName]
    
        # process each row of spreadsheet needing delivery (skipping
        # first row)
        firstRow = True
        for i,c in enumerate(ws.rows):
            if firstRow:
                firstRow = False
                continue
            if i == nRowsMax:
                break
            
            # check if delivery is needed
            if c[10].value.lower()[0] != 'y':
                continue
            # check if any values are None (blank cells in spreadsheet)
            firstName = 'MISSING' if isNone(c[0].value) else c[0].value
            lastName = 'MISSING' if isNone(c[1].value) else c[1].value
            address = 'MISSING' if isNone(c[2].value) else c[2].value
            city = 'MISSING' if isNone(c[3].value) else c[3].value
            state = 'MISSING' if isNone(c[4].value) else c[4].value
            zipCode = '#####' if isNone(c[5].value) else str(c[5].value)
            phone = '(###) ###-####' if isNone(c[6].value) else c[6].value
            phone2 = '(###) ###-####' if isNone(c[7].value) else c[7].value
            email = 'MISSING' if isNone(c[8].value) else c[8].value
            nBags = 0 if isNone(c[9].value) else int(c[9].value)
            notes = '-' if isNone(c[12].value) else c[12].value
        
            # create object and add to outList
            outList.append(MulchContact_t(firstName, lastName, email,
                phone, address, city, state, zipCode, nBags, notes,
                phone2))
        return outList
        # end of method MulchContact_t/read_xlsx
    
    @staticmethod
    def write_kml(mulchList, kmlFileName):
        '''
        Writes .kml file for a list of MulchContact_t objects. This is
        called after address geocoding is complete.
        INPUT
        - mulchList = list of MulchContact objects
        - kmlFileName = name of .kml file to write to
        OUTPUT
        Nothing is returned, a .kml file with the name <kmlFileName> is
        generated.
        '''
        markerFolder = KML.Folder(KML.name("Markers"))
        markerList = []
        for i,c in enumerate(mulchList):
            # check for valid lat/lon
            if any(iter([c.lat < -90.0, c.lat > 90.0, c.lon < -180.0,
                    c.lon > 180.0])):
                continue
            markerList.append(KML.Placemark(
                KML.name(str(c.nBags)),
                KML.description(f'{c.printAddressLabel()}\n' \
                    + f'phone 1: {c.phoneNumber}\n' \
                    + f'phone 2: {c.phoneNumber2}\n' \
                    + f'e-mail: {c.emailAddress}\nNotes: {c.notes}'),
                KML.Point(KML.coordinates(f'{c.lon},{c.lat}'))))
            markerFolder.append(markerList[i])
    
        # write to file
        docKML = KML.Document(KML.name(kmlFileName))
        docKML.append(markerFolder)
        with open(kmlFileName, 'wb') as fout:
            fout.write(etree.tostring(docKML,pretty_print=True))
        # end of method MulchContact_t/write_kml
        
    def printAddressLabel(self):
        '''
        Returns the mailing address as it would appear on an address label.
        '''
        return f'{self.firstName} {self.lastName} ({str(self.nBags)})\n' \
            + f'{self.address}\n{self.city}, {self.state} {self.zipCode}'


def geocode_address(contactList, googleMapsKey):
    '''
    For each Contact_t element in contactList, geocodes address and adds
    address lat/lon to each contactList element.
    INPUT
    - contactList = list of Contact_t objects
    - googleMapsKey = key associated with a Google Cloud account (string)
    OUTPUT
    Nothing is returned. Values are assigned to the .lat and .lon
    attributes of each Contact_t object in contactList.
    '''
    gmaps = googlemaps.Client(key=googleMapsKey)
    
    for c in contactList:
        gr = gmaps.geocode(c.printAddressOneLine())
        if len(gr) < 1:
            msg = 'Bad lat/lon (len(gr)==0) for ' \
                + f'{c.firstName} {c.lastName}'
            print(msg)
            sleep(1)
            continue
        returnedLat = gr[0]['geometry']['location']['lat']
        returnedLon = gr[0]['geometry']['location']['lng']
        if (not isinstance(returnedLat, float))\
            or (not isinstance(returnedLon, float)):
            msg = f'Bad lat/lon for {c.firstName} {c.lastName}'
            print(msg)
            sleep(1)
            continue
        c.lat = returnedLat
        c.lon = returnedLon
        sleep(1)
    # end of function geocode_address

        
def main():
    # Command-line help messages
    descriptionMsg = 'Reads addresses from an Excel file, geocodes each ' \
        + 'address, and puts all the addresses into a KML file for ' \
        + 'viewing in Google Earth.'
    helpDict = {
        'ExcelFile': 'path to .xlsx spreadsheet file',
        'sheetName': 'name of sheet to use in the .xlsx file',
        'nRowsMax': 'max number of rows in .xlsx file to use',
        'GMkey': 'Google Maps API key (string)',
        'kmlFile': 'specify a name for the output KML file. The ' \
            + 'default name is the name of the Excel spreadsheet ' \
            + 'file with the .xlsx file extension replaced with .kml'}
    
    # Parse command-line input
    parser = argparse.ArgumentParser(description=descriptionMsg)
    parser.add_argument('ExcelFile', help=helpDict['ExcelFile'])
    parser.add_argument('sheetName', help=helpDict['sheetName'])
    parser.add_argument('nRowsMax', type=int, help=helpDict['nRowsMax'])
    parser.add_argument('GMkey', help=helpDict['GMkey'])
    parser.add_argument('--kml-file', help=helpDict['kmlFile'])
    args = parser.parse_args()

    # Determine which mode we are using (mulch, flocking, etc...)
    # For now, assume we are using mulch customers
    customerType = MulchContact_t
    
    # read xlsx file
    contactList = customerType.read_xlsx(args.ExcelFile, args.sheetName,
        args.nRowsMax)
    
    # geocode addresses
    geocode_address(contactList, args.GMkey)
    
    # write kml file
    if args.kml_file is None:
        # Use the default name
        outFile = args.ExcelFile.replace('.xlsx', '.kml')
    else:
        # Use user-defined name
        outFile = args.kml_file
    customerType.write_kml(contactList, outFile)
    # end of function main


if __name__ == "__main__":
    main()
